#!/usr/bin/env python3
"""Create the two Italian decision documents for the TESSERA indicator work."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"


def setup(title: str, subtitle: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.3); section.right_margin = Cm(2.3)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]; style.font.name = "Aptos Display"; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(16); style.paragraph_format.space_after = Pt(7)
    p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.add_run(title)
    p = doc.add_paragraph(); p.add_run(subtitle).italic = True
    return doc


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def paragraph(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row): cell.text = value
    for row in t.rows:
        row_properties = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        row_properties.append(no_split)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)


def indicator_note() -> Document:
    doc = setup("BII e MSA GLOBIO 4 per BioMAP", "Nota metodologica per la stima supervisionata con TESSERA v1.1 MPC")
    paragraph(doc, "Questo documento definisce il perimetro scientifico dei due use case. L'obiettivo non è dichiarare che il satellite misuri direttamente la biodiversità o l'abbondanza delle specie: TESSERA apprenderà una relazione supervisionata con mappe ufficiali target, con fonti, scala e incertezza sempre visibili.")
    heading(doc, "Biodiversity Intactness Index")
    paragraph(doc, "Il Biodiversity Intactness Index esprime la condizione media della biodiversità rispetto a uno stato di riferimento. È normalmente rappresentato fra 0 e 1: valori vicini a 1 indicano una comunità più simile alla condizione di riferimento; valori inferiori indicano una maggiore alterazione. La sua interpretazione dipende dalla baseline, dai gruppi biologici e dai dati impiegati nella mappa sorgente.")
    heading(doc, "Mean Species Abundance GLOBIO 4")
    paragraph(doc, "La Mean Species Abundance esprime l'abbondanza media delle specie native rispetto a un ambiente non disturbato. GLOBIO 4 la ricava attraverso relazioni pressione-risposta e non da un singolo canale satellitare. Uso del suolo, frammentazione, infrastrutture, deposizione di azoto e altri driver possono entrare nella definizione a seconda della versione del prodotto.")
    heading(doc, "Implicazione per TESSERA")
    table(doc, ["Aspetto", "Regola operativa"], [["Target", "Un raster BII o MSA ufficiale, con anno, licenza, versione e baseline registrati."], ["Input", "Serie Sentinel-1 e Sentinel-2 elaborate dal checkpoint TESSERA v1.1 MPC."], ["Scala", "Gli embedding a 10 m sono aggregati alla stessa unità fisica del target; non si duplicano etichette grossolane sui pixel."], ["Output", "Stima supervisionata, mappa di validità, incertezza e riferimento alla mappa sorgente."], ["Claim vietato", "Non definire l'output come ricalcolo canonico indipendente di BII o MSA."]])
    heading(doc, "Validazione")
    paragraph(doc, "La validazione usa blocchi geografici indipendenti e quattro macro-regioni europee: Mediterraneo e Italia, Penisola Iberica, Europa centrale e Fennoscandia. Le metriche minime sono MAE, RMSE, bias, R2, calibrazione e intervallo bootstrap del MAE. Il confronto avviene sugli stessi split per baseline Sentinel, embedding congelati e fine-tuning.")
    heading(doc, "Uso decisionale")
    paragraph(doc, "Un utente territoriale potrà usare il risultato per localizzare aree che meritano approfondimento, valutare differenze tra aree e supportare la priorità di un intervento. Non potrà usarlo da solo per attribuire causalità, certificare un miglioramento o sostituire rilievi ecologici indipendenti.")
    heading(doc, "Fonti da fissare nel manifesto")
    paragraph(doc, "TESSERA Temporal Embeddings of Surface Spectra for Earth Representation and Analysis; repository ufficiale ucam-eo/tessera; GeoTessera; fonte ufficiale del raster BII; documentazione e raster ufficiale GLOBIO 4 MSA. Le esatte versioni e licenze saranno riportate nel manifesto della run, non dedotte dal nome dell'indicatore.")
    return doc


def operating_plan() -> Document:
    doc = setup("Piano operativo e risultati TESSERA BioMAP", "Fine tuning BII e MSA GLOBIO 4 con consegna 15 dicembre 2026")
    paragraph(doc, "La consegna è un modulo di ricerca riproducibile e separato dalla pipeline BioAnalyst esistente. Produce due modelli TESSERA fine-tuned, due catene di valutazione e output strutturati pronti per una futura integrazione nella dashboard BioMAP.")
    heading(doc, "Scelta tecnica")
    paragraph(doc, "Il percorso critico usa TESSERA v1.1 MPC nel repository ufficiale ucam-eo/tessera, directory tessera_infer_QAT e checkpoint full MPC. Questa scelta conserva la compatibilità con gli embedding europei GeoTessera e permette di aggiornare realmente i pesi. TESSERA v2 Medium rimane un benchmark successivo, non una dipendenza della consegna.")
    heading(doc, "Calendario")
    table(doc, ["Periodo", "Output verificabile"], [["17-30 settembre", "Raster ufficiali congelati, preflight GPU/checkpoint/dati, manifesto di provenienza."], ["1-16 ottobre", "Campioni armonizzati, split a blocchi, baseline Sentinel e embedding congelati."], ["17-31 ottobre", "Gate di validazione e risultati downstream per BII e MSA."], ["1-20 novembre", "Due fine-tuning v1.1 controllati, checkpoint e log."], ["21 novembre-5 dicembre", "Mappe, maschere, copertura, incertezza e aggregati europei/AOI."], ["6-15 dicembre", "Test finale, documentazione, codice e pacchetto di consegna."]])
    heading(doc, "Artefatti implementati")
    paragraph(doc, "Il modulo tessera_indicators contiene: configurazione unica dei target; preflight che blocca fonti o checkpoint incompleti; split spaziali europei; baseline Ridge su embedding congelati; metriche e intervalli; runner CUDA di fine-tuning v1.1 con pooling per supporto del target; report_context per la futura funzione Map-to-Text.")
    doc.add_page_break()
    heading(doc, "Contratto BioMAP")
    table(doc, ["Campo", "Scopo"], [["selected_area e period", "Territorio e data richiesti dall'utente."], ["indicator e indicator_value", "BII o MSA stimato, con unità e significato dichiarati."], ["uncertainty e valid_coverage", "Qualità del risultato, senza occultare nodata o extrapolazioni."], ["raster e model_version", "Mappa riproducibile e checkpoint usato."], ["sources e limitations", "Provenienza, baseline e limiti per una comunicazione verificabile."]])
    heading(doc, "Criterio di successo")
    paragraph(doc, "Un fine-tuning è accettato solo se migliora, oppure documenta un compromesso chiaramente motivato, rispetto agli embedding congelati sul test spaziale mai usato. I risultati devono poter essere riprodotti con manifesto, configurazione, checkpoint e dati di input; BII e MSA devono avere artefatti e limiti distinti.")
    heading(doc, "Prospettiva di prodotto")
    paragraph(doc, "Il risultato prepara BioMAP a rispondere a una richiesta territoriale con mappa, valore aggregato, incertezza e testo basato su dati strutturati. L'interfaccia e l'eventuale componente Map-to-Text non sono prerequisiti per il 15 dicembre: il contratto dati evita però di dover ricostruire la pipeline quando sarà il momento di esporla agli utenti.")
    return doc


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    indicator_note().save(OUT / "Nota_metodologica_BII_MSA_GLOBIO4_BioMAP.docx")
    operating_plan().save(OUT / "Piano_operativo_e_risultati_TESSERA_BioMAP.docx")


if __name__ == "__main__":
    main()
